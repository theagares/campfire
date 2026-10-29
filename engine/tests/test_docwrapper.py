"""가린 글을 DOCX 로 감쌀 때 XML 에 넣을 수 없는 글자가 있어도 실패하지 않는다.

맥 v0.2.60 M-8: Perplexity 에 붙여 넣은 글(확장이 "upload" 파일로 받음)에 제어 문자가 섞여, docwrapper 의 add_run 에서
lxml 이 "All strings must be XML compatible" 로 막아 /jobs 가 15번 500 이 됐다. PDF 추출 글의 \f(쪽 나눔)도 같다.
"""

from __future__ import annotations

import io

from app.core.masker import docwrapper


def _text(docx_bytes: bytes) -> str:
    from docx import Document

    return "\n".join(p.text for p in Document(io.BytesIO(docx_bytes)).paragraphs)


def test_제어_문자가_섞여도_감싼다():
    masked = "담당자: [이름 마스킹]\x00\x0b 연락처\x0c\n둘째 줄\x1f 끝\ud800￾� 😀\t탭"
    out = docwrapper.wrap_masked_file(masked, "upload")
    assert out["file_name"] == "upload_masked.docx" and out["mime_type"] == docwrapper.DOCX_MIME
    text = _text(out["bytes"])
    assert text == "담당자: [이름 마스킹] 연락처\n둘째 줄 끝� 😀\t탭"


def test_평범한_글은_그대로다():
    out = docwrapper.wrap_masked_file("첫 줄\n\n셋째 줄 [전화번호 마스킹]", "보고서.pdf")
    assert out["file_name"] == "보고서_masked.docx"
    assert _text(out["bytes"]) == "첫 줄\n\n셋째 줄 [전화번호 마스킹]"
