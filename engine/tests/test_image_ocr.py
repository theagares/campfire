"""이미지 속 글자 검사 — 모델 없이: 파서 표시, 표시 바꾸기(위치), 거르기·미검사, GPU 없음."""

from __future__ import annotations

import asyncio
import io

import pytest
from PIL import Image

from app import config
from app.core.detectors import image_ocr
from app.core.parser import STATUS_OK, STATUS_UNSUPPORTED, parse_document
from app.core.parser.docx import extract_docx
from app.core.parser import images as image_marks
from app.core.parser.images import MARK_RE
from app.core.parser.pptx import extract_pptx


def _png(w: int = 400, h: int = 120, color=(255, 255, 255)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


def _docx_with_picture() -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("첫 문단")
    d.add_picture(io.BytesIO(_png()))
    d.add_paragraph("끝 문단")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx_with_picture() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(1))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_기능을_끄면_파서_결과가_그대로고_이미지_파일은_미지원():
    b = _docx_with_picture()
    assert extract_docx(b) == extract_docx(b, None) == "첫 문단\n\n끝 문단"
    assert parse_document(_png(), "image/png", "shot.png")[1] == STATUS_UNSUPPORTED


def test_꺼져_있으면_문서_속_그림은_세고_표시만_지워_글은_지금과_같다():
    # 이미지 속 글자 검사가 꺼져 있어도 그림을 센다 — orchestrator 가 imagesRemoved 로 원본 대신 글만 보내게 한다.
    data = _docx_with_picture()
    plain, _, _ = parse_document(data, "", "a.docx")
    images: list[bytes] = []
    marked, status, _ = parse_document(data, "", "a.docx", images)
    assert status == STATUS_OK and len(images) == 1
    assert image_marks.strip_marks(marked) == plain


def test_DOCX_표_칸_안의_그림도_그_자리에_표시가_남는다():
    from docx import Document

    d = Document()
    d.add_paragraph("앞 문단")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "캡처"
    t.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(_png()))
    d.add_paragraph("뒤 문단")
    buf = io.BytesIO()
    d.save(buf)
    images: list[bytes] = []
    parts = extract_docx(buf.getvalue(), images).split("\n\n")
    assert len(images) == 1 and parts[0] == "앞 문단" and parts[2] == "뒤 문단"
    assert parts[1].startswith("캡처 | ") and MARK_RE.search(parts[1])


def test_켜면_이미지_파일과_DOCX_PPTX_속_그림_자리에_표시가_남는다():
    images: list[bytes] = []
    text, status, _ = parse_document(_png(), "image/png", "shot.png", images)
    assert status == STATUS_OK and MARK_RE.fullmatch(text) and len(images) == 1

    images = []
    text = extract_docx(_docx_with_picture(), images)
    assert len(images) == 1 and text.split("\n\n")[0] == "첫 문단" and MARK_RE.search(text) and text.endswith("끝 문단")

    images = []
    assert MARK_RE.search(extract_pptx(_pptx_with_picture(), images)) and len(images) == 1


@pytest.fixture(autouse=True)
def _no_commit_limit(monkeypatch):
    """이 PC 의 실제 남은 시스템 커밋에 시험이 흔들리지 않게 기본은 '모름'(메모리 안전장치 시험만 따로 흉내 낸다)."""
    monkeypatch.setattr(image_ocr, "free_commit_mb", lambda: None)
    monkeypatch.setattr(image_ocr, "free_physical_mb", lambda: None)
    monkeypatch.setattr(image_ocr, "commit_limit_mb", lambda: None)


_real_gpu_worker_first = image_ocr.gpu_worker_first  # 아래 autouse 가 바꾸기 전 — 순서 시험이 진짜를 쓴다


@pytest.fixture(autouse=True)
def _sidecar_by_default(monkeypatch):
    """기존 시험은 사이드카 경로를 본다 — NVIDIA PC 에서 시험이 진짜 통합 worker(모델)로 OCR 을 보내지 않게(worker 안 OCR
    시험은 가짜 worker 를 따로 끼운다)."""
    monkeypatch.setattr(config, "IMAGE_OCR_IN_WORKER", False)


@pytest.fixture(autouse=True)
def _no_gpu_worker(monkeypatch):
    """사이드카를 띄우는 시험이 진짜 통합 GPU worker(PII·인젝션 모델)를 올리지 않게 한다(Windows 는 기본 켜짐)."""
    async def nothing():
        return None

    monkeypatch.setattr(image_ocr, "gpu_worker_first", nothing)


@pytest.fixture(autouse=True)
def _empty_ocr_cache():
    """업로드 사이 OCR 캐시는 프로세스 전역이다 — 같은 _png() 바이트를 쓰는 시험끼리 결과가 새지 않게 비운다."""
    image_ocr._cache.clear()
    yield
    image_ocr._cache.clear()


def _fake_reader(monkeypatch, results):
    calls = []

    async def read(blob):
        calls.append(blob)
        return results.pop(0)

    monkeypatch.setattr(image_ocr.get(), "read", read)
    return calls


def test_표시를_읽은_글로_바꾸고_이미지_위치를_바꾼_뒤_글_기준으로_적는다(monkeypatch):
    a, b = _png(color=(250, 250, 250)), _png(color=(240, 240, 240))
    _fake_reader(monkeypatch, [
        {"text": "가나다", "tiles": [{"box": [0, 0, 10, 10], "start": 0, "end": 3}], "secs": 1},
        {"text": "주민 900101-1234567", "tiles": [{"box": [5, 5, 50, 20], "start": 3, "end": 17}], "secs": 1},
    ])
    images = [a, b]
    text, infos = asyncio.run(image_ocr.fill("앞\n0\n중간\n1\n뒤", images))
    assert [i["status"] for i in infos] == ["read", "read"]
    assert text == "앞\n[이미지 1]\n가나다\n중간\n[이미지 2]\n주민 900101-1234567\n뒤"
    assert text[infos[0]["start"]:infos[0]["end"]] == "[이미지 1]\n가나다"
    assert text[infos[1]["start"]:infos[1]["end"]] == "[이미지 2]\n주민 900101-1234567"  # 첫 이미지 글 길이만큼 밀렸다


def test_작은_그림은_건너뛰고_못_읽은_그림은_미검사(monkeypatch):
    _fake_reader(monkeypatch, [None])
    images = [_png(40, 20), _png()]
    text, infos = asyncio.run(image_ocr.fill("0\n1", images))
    assert [i["status"] for i in infos] == ["skipped", "unchecked"]
    assert infos[1]["reason"] == image_ocr.UNCHECKED_REASON
    assert text == "\n"  # 둘 다 글 없이 빠진다


def test_같은_그림은_한_번만_읽는다(monkeypatch):
    calls = _fake_reader(monkeypatch, [{"text": "로고 아님", "tiles": [], "secs": 1}])
    blob = _png()
    text, infos = asyncio.run(image_ocr.fill("0 1", [blob, blob]))
    assert len(calls) == 1 and [i["status"] for i in infos] == ["read", "read"] and text.count("로고 아님") == 2


def test_다른_업로드에서_읽은_그림은_30분_안에_다시_읽지_않는다(monkeypatch):
    # 맥 v0.2.60: 그림 든 문서가 OCR 을 기다리다 Claude 업로드 시간이 지났다 — 다시 붙이면 곧바로 끝나야 한다.
    calls = _fake_reader(monkeypatch, [{"text": "고객 김민준", "tiles": [], "secs": 7.5}])
    blob = _png()
    first, _ = asyncio.run(image_ocr.fill("0", [blob]))
    again, infos = asyncio.run(image_ocr.fill("앞 0", [blob]))  # 다른 문서·다른 업로드
    assert len(calls) == 1, "다른 업로드의 같은 그림을 다시 읽었다"
    assert first == "[이미지 1]\n고객 김민준" and again == "앞 [이미지 1]\n고객 김민준"
    assert infos[0]["status"] == "read" and infos[0]["secs"] == 0.0


def test_못_읽은_그림은_담지_않아_다음에_다시_읽는다(monkeypatch):
    calls = _fake_reader(monkeypatch, [None, {"text": "다시 읽음", "tiles": [], "secs": 1}])
    blob = _png()
    _, infos = asyncio.run(image_ocr.fill("0", [blob]))
    assert infos[0]["status"] == "unchecked"
    text, infos = asyncio.run(image_ocr.fill("0", [blob]))
    assert len(calls) == 2 and infos[0]["status"] == "read" and text.endswith("다시 읽음")


def test_OCR_캐시는_가득_차면_오래된_것부터_밀어낸다(monkeypatch):
    a, b = _png(color=(250, 250, 250)), _png(color=(240, 240, 240))
    calls = _fake_reader(monkeypatch, [{"text": t, "tiles": [], "secs": 1} for t in ("가", "나", "가")])
    monkeypatch.setattr(image_ocr, "_CACHE_MAX", 1)
    for blob in (a, b, a):  # b 가 들어오며 a 가 밀려났다
        asyncio.run(image_ocr.fill("0", [blob]))
    assert len(calls) == 3


def test_OCR_캐시는_30분이_지나면_다시_읽는다(monkeypatch):
    calls = _fake_reader(monkeypatch, [{"text": "가", "tiles": [], "secs": 1}, {"text": "가", "tiles": [], "secs": 1}])
    monkeypatch.setattr(image_ocr, "_CACHE_TTL_SEC", -1.0)  # 담자마자 지난 것으로
    blob = _png()
    asyncio.run(image_ocr.fill("0", [blob]))
    asyncio.run(image_ocr.fill("0", [blob]))
    assert len(calls) == 2, "시간이 지난 OCR 결과를 다시 썼다"


def _varco_runtime():
    """사이드카 런타임(엔진 패키지 밖 스크립트)을 모듈로 불러온다 — torch·transformers 가 있어야 한다."""
    import importlib.util
    from pathlib import Path

    pytest.importorskip("torch")
    pytest.importorskip("transformers")
    runtime = Path(__file__).resolve().parents[1] / "app" / "models" / "varco_ocr" / "runtime" / "varco_ocr_inference.py"
    spec = importlib.util.spec_from_file_location("varco_rt", runtime)
    rt = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rt)
    return rt


def test_VARCO_출력의_로고_수식_조각은_버리고_동그라미_숫자는_살린다():
    """회사소개서 로고가 '\\circle \\textcircled{1} ㅁ | 르온테 크' 로 남았다(서버 실측)."""
    clean = _varco_runtime().clean
    assert clean("\\circle") == ""
    assert clean("\\textcircled{1} 주문 자동 수집\n\\textcircled{2} 실시간 재고") == "① 주문 자동 수집\n② 실시간 재고"
    assert clean("\\textcircled{1}") == "①"
    assert clean("\\textcircled{0} 라이브러리") == "⓪ 라이브러리"
    assert clean("\\red{선} \\blue{원} \\orange{선}") == "선 원 선"  # 색만 입힌 범례 글자는 살린다
    assert clean("\\textcolor{red}{주의} 사항") == "주의 사항"
    assert clean("\\begin{tikzpicture}") == ""  # 그림 코드는 여전히 거른다
    assert clean("\\circle{blue} \\hline INJECTION 모델 \\mid LG EXAONE") == "INJECTION 모델 | LG EXAONE"  # 아이콘 도형은 글자가 아니다
    assert clean("\\textsuperscript{V}RAM 4GB") == "VRAM 4GB"
    assert clean("경로 C:\\Users\\홍길동") == "경로 C:\\Users\\홍길동"  # 아무 '\\단어' 나 지우지 않는다
    assert clean("\\hookrightarrow 정책 종료") == "↪ 정책 종료"


def test_VARCO_는_같은_줄을_되풀이하면_상한까지_가지_않고_멈춘다():
    """로고 조각이 "VARCO-VISION\\n"(6토큰)을 2048토큰까지 뽑아 그림 한 장이 77초 걸렸다(서버 실측)."""
    rt = _varco_runtime()
    import torch

    n_in = 10
    stop = rt._stopper(torch.tensor([2048]), n_in)[0]
    line = [11, 12, 13, 14, 15, 99]  # 한 줄 + 줄바꿈
    assert stop(torch.tensor([[1] * n_in + line * 17]), None).item()               # 96토큰 넘게 같은 줄 → 멈춤
    assert not stop(torch.tensor([[1] * n_in + line * 5 + list(range(100, 170))]), None).item()  # 다른 글이 이어지면 계속
    assert not stop(torch.tensor([[1] * n_in + list(range(100, 230))]), None).item()


def test_GPU_가_없다고_한_번_보면_다시_띄우지_않는다(monkeypatch):
    proc = image_ocr.VarcoOcrProcess()
    monkeypatch.setattr(image_ocr.model_status, "image_ocr_ready", lambda: True)
    spawned = []

    async def ensure():
        spawned.append(1)
        raise RuntimeError("image_ocr: 예상치 못한 준비 응답: {'ready': False, 'error': 'no_gpu'}")

    monkeypatch.setattr(proc, "_ensure_process", ensure)
    assert asyncio.run(proc.read(_png())) is None and proc.no_gpu
    assert asyncio.run(proc.read(_png())) is None and spawned == [1]


def test_사이드카를_못_띄우면_로그를_남기고_GPU_없음이_아니라_그_이유로_막는다(monkeypatch, caplog):
    # v0.2.56 Windows 실측(EN-6): 메모리 부족으로 사이드카가 죽었는데 "GPU 실행 환경이 없거나…" 로 막고 로그도 없었다.
    proc = image_ocr.VarcoOcrProcess()
    monkeypatch.setattr(image_ocr, "_process", proc)
    monkeypatch.setattr(image_ocr.model_status, "image_ocr_ready", lambda: True)

    async def ensure():
        raise RuntimeError("image_ocr: 서브프로세스가 준비 전 종료됨: OSError: [WinError 1455] 페이징 파일이 너무 작습니다")

    monkeypatch.setattr(proc, "_ensure_process", ensure)
    with caplog.at_level("WARNING", logger="securedoc.engine.image_ocr"):
        _text, infos = asyncio.run(image_ocr.fill("0", [_png()]))
    assert infos[0]["status"] == "unchecked"
    assert infos[0]["reason"] == image_ocr.LAUNCH_FAILED_REASON
    assert not proc.no_gpu  # GPU 가 없는 게 아니다 — 다음 그림에서 다시 띄워 본다
    assert "WinError 1455" in caplog.text


def _guarded_proc(monkeypatch, *, free_mb, running):
    """남은 시스템 커밋을 free_mb 로 흉내 낸 사이드카 — 띄우기·읽기·내리기를 기록한다."""
    proc = image_ocr.VarcoOcrProcess()
    proc._process_ready = running
    calls = {"launch": 0, "rpc": 0, "kill": 0}
    monkeypatch.setattr(image_ocr.model_status, "image_ocr_ready", lambda: True)
    monkeypatch.setattr(image_ocr, "free_commit_mb", lambda: free_mb)

    async def ensure():
        calls["launch"] += 1

    async def rpc(_payload):
        calls["rpc"] += 1
        return {"results": [{"text": "읽음", "tiles": [], "secs": 1.0}]}

    async def kill():
        calls["kill"] += 1

    monkeypatch.setattr(proc, "_ensure_process", ensure)
    monkeypatch.setattr(proc, "_rpc", rpc)
    monkeypatch.setattr(proc, "_kill_process", kill)
    return proc, calls


def test_남은_시스템_메모리가_모자라면_사이드카를_띄우지_않고_미검사(monkeypatch):
    # v0.2.62 측정(RAM 15GB): 사이드카 커밋 8GB 로 PC 남은 커밋이 10.4GB → 0.67GB — 띄우기 전에 막는다.
    proc, calls = _guarded_proc(monkeypatch, free_mb=config.IMAGE_OCR_LAUNCH_FREE_COMMIT_MB - 1, running=False)
    assert asyncio.run(proc.read(_png())) is None
    assert proc.last_failure == image_ocr.LOW_MEMORY_REASON and calls == {"launch": 0, "rpc": 0, "kill": 0}


def test_떠_있을_때_메모리가_바닥나면_읽지_않고_내린다(monkeypatch):
    proc, calls = _guarded_proc(monkeypatch, free_mb=config.IMAGE_OCR_READ_FREE_COMMIT_MB - 1, running=True)
    assert asyncio.run(proc.read(_png())) is None
    assert proc.last_failure == image_ocr.LOW_MEMORY_REASON
    assert calls == {"launch": 0, "rpc": 0, "kill": 1} and not proc._process_ready, "메모리를 돌려주려면 사이드카를 내려야 한다"


def test_커밋이_넉넉해도_남은_물리_메모리가_모자라면_띄우지_않는다(monkeypatch):
    # v0.2.63 후보 3차(페이지 파일을 늘린 뒤): 남은 커밋 11.9GB 인데 WDDM 이 VRAM 을 RAM 으로 밀어내 물리 0.2GB
    proc, calls = _guarded_proc(monkeypatch, free_mb=40000, running=False)
    monkeypatch.setattr(image_ocr, "free_physical_mb", lambda: config.IMAGE_OCR_LAUNCH_FREE_PHYS_MB - 1)
    assert asyncio.run(proc.read(_png())) is None
    assert proc.last_failure == image_ocr.LOW_MEMORY_REASON and calls == {"launch": 0, "rpc": 0, "kill": 0}


def test_떠_있을_때_물리_메모리가_바닥나면_읽지_않고_내리고_읽기_기준_위면_읽는다(monkeypatch):
    proc, calls = _guarded_proc(monkeypatch, free_mb=40000, running=True)
    monkeypatch.setattr(image_ocr, "free_physical_mb", lambda: config.IMAGE_OCR_READ_FREE_PHYS_MB - 1)
    assert asyncio.run(proc.read(_png())) is None and calls["kill"] == 1 and not proc._process_ready
    proc, calls = _guarded_proc(monkeypatch, free_mb=40000, running=True)
    monkeypatch.setattr(image_ocr, "free_physical_mb", lambda: config.IMAGE_OCR_READ_FREE_PHYS_MB + 1)  # 띄우기 기준보다는 아래
    assert asyncio.run(proc.read(_png()))["text"] == "읽음" and calls["kill"] == 0


def test_사이드카_준비_응답의_묶음_예산_그래프를_로그에_남긴다(monkeypatch, caplog):
    async def ready(_self, _timeout, _tag):
        return {"ready": True, "device": "cuda", "batch": 4, "budget": 4608, "graph": True, "card_free_mb": 2662}

    monkeypatch.setattr(image_ocr.EncoderPiiDetector, "_read_ready_line", ready)
    with caplog.at_level("INFO", logger="securedoc.engine.image_ocr"):
        payload = asyncio.run(image_ocr.VarcoOcrProcess()._read_ready_line(10, "image_ocr"))
    assert payload["batch"] == 4
    assert "'batch': 4" in caplog.text and "'graph': True" in caplog.text and "'card_free_mb': 2662" in caplog.text


def test_메모리가_넉넉하거나_모르면_지금처럼_읽는다(monkeypatch):
    # 떠 있으면 읽기 기준(2GB)만 넘으면 된다 — 띄우기 기준(9GB)보다 적어도 이미 올린 사이드카로 읽는다
    proc, calls = _guarded_proc(monkeypatch, free_mb=config.IMAGE_OCR_READ_FREE_COMMIT_MB + 1000, running=True)
    assert asyncio.run(proc.read(_png()))["text"] == "읽음" and calls["rpc"] == 1
    proc, calls = _guarded_proc(monkeypatch, free_mb=None, running=False)  # Windows 가 아니거나 모름
    assert asyncio.run(proc.read(_png()))["text"] == "읽음" and calls == {"launch": 1, "rpc": 1, "kill": 0}


def test_메모리_부족_까닭을_커밋과_물리로_나눠_health_에_싣고_다시_읽으면_지운다(monkeypatch):
    # 커밋이면 페이지 파일을 늘리면 풀리고(앱이 방법 안내), 물리면 다른 프로그램을 닫으면 풀린다 — 앱이 둘을 달리 안내한다
    proc, _calls = _guarded_proc(monkeypatch, free_mb=5423.4, running=False)  # 1차 실측: worker 뒤 남은 커밋
    monkeypatch.setattr(image_ocr, "_process", proc)
    monkeypatch.setattr(image_ocr, "commit_limit_mb", lambda: 33669.4)
    assert asyncio.run(proc.read(_png())) is None
    short = image_ocr.status()["memoryShort"]
    assert {k: short[k] for k in ("kind", "freeMb", "needMb", "commitLimitMb")} == {
        "kind": "commit", "freeMb": 5423, "needMb": config.IMAGE_OCR_LAUNCH_FREE_COMMIT_MB, "commitLimitMb": 33669}
    assert short["at"].startswith("20")

    monkeypatch.setattr(image_ocr, "free_commit_mb", lambda: 40000)
    monkeypatch.setattr(image_ocr, "free_physical_mb", lambda: config.IMAGE_OCR_LAUNCH_FREE_PHYS_MB - 1)
    assert asyncio.run(proc.read(_png())) is None
    assert image_ocr.status()["memoryShort"]["kind"] == "phys"

    monkeypatch.setattr(image_ocr, "free_physical_mb", lambda: 8000)
    assert asyncio.run(proc.read(_png()))["text"] == "읽음"
    assert image_ocr.status()["memoryShort"] is None, "다시 읽었으면 안내를 거둔다"


def test_윈도우가_아니면_메모리를_모르고_memoryShort_도_생기지_않는다(monkeypatch):
    import types

    with monkeypatch.context() as m:  # 모듈의 os 만 — 전역 os.name 을 바꾸면 pathlib 이 깨진다
        m.setattr(image_ocr, "os", types.SimpleNamespace(name="posix"))
        assert image_ocr._memory_status() is None  # free_commit_mb·free_physical_mb·commit_limit_mb 가 모두 이것을 지난다
    proc, _calls = _guarded_proc(monkeypatch, free_mb=None, running=False)
    monkeypatch.setattr(image_ocr, "_process", proc)
    assert asyncio.run(proc.read(_png()))["text"] == "읽음"
    assert image_ocr.status()["memoryShort"] is None


def test_사이드카를_띄우기_전에_통합_GPU_worker_를_먼저_올린다(monkeypatch):
    # worker 가 유휴로 내려간 뒤 OCR 이 먼저 떠 남은 VRAM 으로 큰 묶음을 잡으면, 뒤에 올라오는 worker(약 4GB)와 12GB 카드가 넘친다.
    order = []
    proc, _calls = _guarded_proc(monkeypatch, free_mb=None, running=False)
    monkeypatch.setattr(image_ocr, "free_commit_mb", lambda: order.append("commit"))
    real_ensure = proc._ensure_process

    async def ensure():
        order.append("launch")
        await real_ensure()

    class _Worker:
        async def ensure_ready(self):
            order.append("worker")

    monkeypatch.setattr(proc, "_ensure_process", ensure)
    monkeypatch.setattr(image_ocr, "gpu_worker_first", _real_gpu_worker_first)
    monkeypatch.setattr(image_ocr.combined_gpu, "get_worker", _Worker)
    assert asyncio.run(proc.read(_png()))["text"] == "읽음"
    assert order == ["worker", "commit", "launch"], "worker 몫을 뺀 남은 메모리로 띄우기 기준·묶음을 정해야 한다"

    order.clear()  # 이미 떠 있으면 worker 를 건드리지 않는다
    proc._process_ready = True
    asyncio.run(proc.read(_png()))
    assert "worker" not in order


def test_통합_worker_를_못_쓰면_OCR_은_그대로_띄운다(monkeypatch):
    proc, calls = _guarded_proc(monkeypatch, free_mb=None, running=False)

    class _Disabled:
        async def ensure_ready(self):
            raise image_ocr.combined_gpu.CombinedGpuUnavailable("combined GPU worker is disabled")

    monkeypatch.setattr(image_ocr, "gpu_worker_first", _real_gpu_worker_first)
    monkeypatch.setattr(image_ocr.combined_gpu, "get_worker", _Disabled)
    assert asyncio.run(proc.read(_png()))["text"] == "읽음" and calls["launch"] == 1


def test_GPU_torch_는_이미지_검사_사이드카에만_앞세우고_따로_지정한_파이썬이면_그대로(monkeypatch, tmp_path):
    monkeypatch.setattr(config, "IMAGE_OCR_GPU_TORCH_DIR", tmp_path / "gpu_torch")
    monkeypatch.setattr(config, "IMAGE_OCR_PYTHON_OVERRIDE", "")
    monkeypatch.setattr(image_ocr.model_status, "image_ocr_accelerator", lambda: "cuda")
    monkeypatch.delenv("PYTHONPATH", raising=False)
    assert image_ocr.VarcoOcrProcess()._spawn_env() == {"PYTHONPATH": str(tmp_path / "gpu_torch")}
    monkeypatch.setattr(config, "IMAGE_OCR_PYTHON_OVERRIDE", r"C:\cuda\python.exe")
    assert image_ocr.VarcoOcrProcess()._spawn_env() == {}


def test_Apple_MPS는_번들_torch를_쓰고_지원안되는_연산만_CPU로_폴백한다(monkeypatch):
    monkeypatch.setattr(config, "IMAGE_OCR_PYTHON_OVERRIDE", "")
    monkeypatch.setattr(image_ocr.model_status, "image_ocr_accelerator", lambda: "mps")
    assert image_ocr.VarcoOcrProcess()._spawn_env() == {"PYTORCH_ENABLE_MPS_FALLBACK": "1"}


@pytest.mark.parametrize("enabled", [False, True])
def test_켰을_때만_이미지_형식이_지원_목록에_있다(monkeypatch, enabled):
    import importlib

    monkeypatch.setenv("SECUREDOC_IMAGE_OCR_ENABLED", "1" if enabled else "0")
    try:
        cfg = importlib.reload(config)
        assert (".png" in cfg.SUPPORTED_EXTENSIONS) is enabled
    finally:
        monkeypatch.delenv("SECUREDOC_IMAGE_OCR_ENABLED", raising=False)
        importlib.reload(config)


GRAPH_FAILED = "image_ocr: 추론 오류: graph_failed: AcceleratorError: CUDA error: an illegal memory access was encountered"
DIED = "image_ocr: 서브프로세스가 응답 없이 종료됨: "  # 표식도 못 남기고 끝났다(그래프 탓이거나 외부 종료·메모리 부족)


def _graph_proc(monkeypatch, *errors: str):
    """CUDA 그래프를 켠 사이드카 — 처음 읽기들이 errors 로 차례로 실패하고 그 뒤로는 읽는다. 띄우기·내리기·요청을 기록한다."""
    monkeypatch.setattr(config, "IMAGE_OCR_CUDA_GRAPH", True)
    proc = image_ocr.VarcoOcrProcess()
    monkeypatch.setattr(image_ocr.model_status, "image_ocr_ready", lambda: True)
    calls = {"spawn": [], "kill": 0, "rpc": 0}

    async def ensure():
        if not proc._process_ready:
            calls["spawn"].append(proc._spawn_args()[-1])  # --cuda-graph 값
            proc._process_ready = True

    async def kill():
        calls["kill"] += 1

    async def rpc(_payload):
        calls["rpc"] += 1
        if calls["rpc"] <= len(errors):
            raise RuntimeError(errors[calls["rpc"] - 1])
        return {"results": [{"text": "읽음", "tiles": [], "secs": 1.0}]}

    monkeypatch.setattr(proc, "_ensure_process", ensure)
    monkeypatch.setattr(proc, "_kill_process", kill)
    monkeypatch.setattr(proc, "_rpc", rpc)
    return proc, calls


def test_graph_failed_표식이면_그래프를_끈_사이드카로_다시_띄워_같은_그림을_다시_읽고_이후로도_끈다(monkeypatch, caplog):
    proc, calls = _graph_proc(monkeypatch, GRAPH_FAILED)
    with caplog.at_level("WARNING", logger="securedoc.engine.image_ocr"):
        assert asyncio.run(proc.read(_png()))["text"] == "읽음", "그림이 막히면 안 된다"
    assert calls == {"spawn": ["1", "0"], "kill": 1, "rpc": 2} and not proc.cuda_graph
    assert "CUDA 그래프 경로가 실패해 이번 실행 동안 끈다" in caplog.text
    proc._process_ready = False  # 유휴로 내려간 뒤 다시 띄워도 그래프는 꺼진 채
    assert asyncio.run(proc.read(_png()))["text"] == "읽음"
    assert calls["spawn"] == ["1", "0", "0"] and calls["kill"] == 1


def test_답_없이_한_번_죽으면_그래프를_켠_채_다시_띄워_다시_읽는다(monkeypatch, caplog):
    # v0.2.63 후보 3차: 측정 감시가 메모리 때문에 끈 것을 그래프 실패로 받아들여 그 실행 동안 그래프가 꺼졌다
    proc, calls = _graph_proc(monkeypatch, DIED)
    with caplog.at_level("WARNING", logger="securedoc.engine.image_ocr"):
        assert asyncio.run(proc.read(_png()))["text"] == "읽음"
    assert calls == {"spawn": ["1", "1"], "kill": 1, "rpc": 2} and proc.cuda_graph
    assert "그래프를 켠 채 한 번 다시 띄워" in caplog.text


def test_같은_그림에서_또_답_없이_죽으면_그때_그래프를_끈다(monkeypatch):
    proc, calls = _graph_proc(monkeypatch, DIED, DIED)
    assert asyncio.run(proc.read(_png()))["text"] == "읽음", "그림이 막히면 안 된다"
    assert calls == {"spawn": ["1", "1", "0"], "kill": 2, "rpc": 3} and not proc.cuda_graph
    # 다음 그림은 처음부터 다시 센다 — 그림 하나에서 한 번 죽은 걸로는 끄지 않았다
    proc2, calls2 = _graph_proc(monkeypatch, DIED)
    assert asyncio.run(proc2.read(_png()))["text"] == "읽음" and proc2.cuda_graph


def test_답_없이_죽은_뒤_메모리가_모자라면_다시_띄우지_않고_그림만_뺀다(monkeypatch):
    proc, calls = _graph_proc(monkeypatch, DIED)
    proc._process_ready = True  # 이미 떠 있다

    async def died_then_short(_payload):  # 읽다 죽었고, 그 사이 물리 메모리가 띄우기 기준 아래로
        calls["rpc"] += 1
        monkeypatch.setattr(image_ocr, "free_physical_mb", lambda: config.IMAGE_OCR_LAUNCH_FREE_PHYS_MB - 1)
        raise RuntimeError(DIED)

    monkeypatch.setattr(proc, "_rpc", died_then_short)
    assert asyncio.run(proc.read(_png())) is None
    assert proc.last_failure == image_ocr.LOW_MEMORY_REASON, "막지 말고 그림만 빼야 한다((가))"
    assert calls["spawn"] == [] and calls["rpc"] == 1


def test_그래프와_무관한_오류는_지금처럼_미검사로_다시_띄우지_않는다(monkeypatch):
    proc, calls = _graph_proc(monkeypatch, "image_ocr: 추론 오류: ValueError: 그림을 열 수 없음")
    assert asyncio.run(proc.read(_png())) is None and proc.last_failure == image_ocr.READ_FAILED_REASON
    assert calls == {"spawn": ["1"], "kill": 0, "rpc": 1} and proc.cuda_graph


def test_CUDA_그래프_env_를_끄면_처음부터_끈_사이드카로(monkeypatch):
    monkeypatch.setattr(config, "IMAGE_OCR_CUDA_GRAPH", False)
    assert image_ocr.VarcoOcrProcess()._spawn_args()[-2:] == ["--cuda-graph", "0"]
