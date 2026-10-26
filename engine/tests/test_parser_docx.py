"""DOCX 파서 — 표가 본문에서 제자리에 들어가는지, 병합된 칸이 한 번만 나오는지."""
from __future__ import annotations

import io

from app.core.parser.docx import extract_docx


def _docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("1. 고객 정보")
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).text = "이름"
    t.cell(0, 1).text = "연락처"
    t.cell(0, 2).text = "비고"
    t.cell(1, 0).text = "홍길동"
    t.cell(1, 1).merge(t.cell(1, 2)).text = "010-1234-5678"  # 가로로 병합된 칸
    d.add_paragraph("2. 처리 결과")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_표는_앞뒤_문단_사이_제자리에_들어가고_병합된_칸은_한_번만_나온다():
    assert extract_docx(_docx()).split("\n\n") == [
        "1. 고객 정보",
        "이름 | 연락처 | 비고",
        "홍길동 | 010-1234-5678",
        "2. 처리 결과",
    ]
