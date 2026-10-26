"""좌표 기반 문서 파서 — 표/글/빈 셀이 원래 읽기 위치를 유지하는지 검증."""
from __future__ import annotations

import io

from app.core.parser.docx import extract_docx
from app.core.parser.pdf import extract_pdf
from app.core.parser.pptx import extract_pptx
from app.core.parser.xlsx import extract_xlsx


def _pdf_with_middle_table() -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 52), "TOP TEXT")
    for y in (92, 132, 172):
        page.draw_line((72, y), (300, y))
    for x in (72, 186, 300):
        page.draw_line((x, 92), (x, 172))
    page.insert_text((82, 117), "H1")
    page.insert_text((196, 117), "H2")
    page.insert_text((82, 157), "A1")
    page.insert_text((196, 157), "A2")
    page.insert_text((72, 232), "BOTTOM TEXT")
    data = doc.tobytes()
    doc.close()
    return data


def test_pdf_inserts_table_at_its_visual_position_without_duplicate_text():
    assert extract_pdf(_pdf_with_middle_table()).splitlines() == [
        "TOP TEXT",
        "H1 | H2",
        "A1 | A2",
        "BOTTOM TEXT",
    ]


def test_pptx_uses_visual_position_not_shape_creation_order_and_reads_groups():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # 생성 순서는 일부러 화면 순서의 반대로 둔다.
    slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(0.5)).text = "BOTTOM"
    slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(4), Inches(0.5)).text = "TOP"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "LEFT"
    table.cell(0, 1).text = "RIGHT"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(0.5)).text = "GROUP"

    buf = io.BytesIO()
    prs.save(buf)
    assert extract_pptx(buf.getvalue()).splitlines() == [
        "[슬라이드 1]",
        "TOP",
        "LEFT | RIGHT",
        "GROUP",
        "BOTTOM",
    ]


def test_docx_keeps_nested_table_between_cell_paragraphs_and_vertical_merge_column():
    from docx import Document

    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].text = "CELL BEFORE"
    nested = cell.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "NEST A"
    nested.cell(0, 1).text = "NEST B"
    cell.add_paragraph("CELL AFTER")

    merged = doc.add_table(rows=2, cols=2)
    merged.cell(0, 0).text = "MERGED"
    merged.cell(0, 0).merge(merged.cell(1, 0))
    merged.cell(0, 1).text = "ROW 1"
    merged.cell(1, 1).text = "ROW 2"

    buf = io.BytesIO()
    doc.save(buf)
    assert extract_docx(buf.getvalue()).split("\n\n") == [
        "CELL BEFORE\nNEST A | NEST B\nCELL AFTER",
        "MERGED | ROW 1",
        " | ROW 2",
    ]


def test_xlsx_keeps_empty_columns_between_values():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet["A1"] = "A"
    sheet["C1"] = "C"
    sheet["B2"] = "B"
    buf = io.BytesIO()
    workbook.save(buf)
    workbook.close()

    assert extract_xlsx(buf.getvalue()).splitlines() == ["[Sparse]", "A |  | C", " | B"]
